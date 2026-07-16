from http.server import BaseHTTPRequestHandler
import json
import io
import openpyxl


REV_MAP = {
    "A": "P00", "B": "P01", "C": "P02", "D": "P03",
    "E": "P04", "F": "P05", "G": "P06", "H": "P07",
    "I": "P08", "J": "P09", "K": "P10", "L": "P11"
}

def convert_rev(rev):
    return REV_MAP.get(rev.upper().strip(), rev.strip())


def read_raw_bom(wb):
    ws = wb.worksheets[0]
    rows = []
    for r in range(2, ws.max_row + 1):
        r0 = ws.cell(r, 1).value
        if r0 is None: continue
        try: level = int(float(str(r0)))
        except: continue
        title    = str(ws.cell(r, 2).value or "").strip()
        rev      = str(ws.cell(r, 3).value or "").strip()
        desc     = str(ws.cell(r, 4).value or "").strip()
        maturity = str(ws.cell(r, 5).value or "").strip()
        if not title: continue
        rows.append({"level": level, "title": title, "rev": rev,
                     "desc": desc, "maturity": maturity, "row": r})
    return rows


def find_parent_title(rows, current_idx, current_level):
    for i in range(current_idx - 1, -1, -1):
        if rows[i]["level"] == current_level - 1:
            return rows[i]["title"]
    return "Root"


def build_quantified_bom(rows, convert_revision=False):
    # Her satir icin parent hesapla
    for i, row in enumerate(rows):
        row["parent"] = find_parent_title(rows, i, row["level"])
        row["status"] = "PENDING"

    result = []

    for i, row in enumerate(rows):
        if row["status"] in ("DONE", "SKIP"):
            continue

        level   = row["level"]
        title   = row["title"]
        parent  = row["parent"]
        qty     = 1
        row["status"] = "DONE"

        # Duplicate bul
        for j in range(i + 1, len(rows)):
            if rows[j]["status"] in ("DONE", "SKIP"):
                continue
            if (rows[j]["level"] == level and
                rows[j]["title"] == title and
                rows[j]["parent"] == parent):
                qty += 1
                rows[j]["status"] = "SKIP"

        # Rev donustur
        rev = convert_rev(row["rev"]) if convert_revision else row["rev"]

        result.append({
            "level":    level,
            "title":    title,
            "rev":      rev,
            "desc":     row["desc"],
            "maturity": row["maturity"],
            "parent":   parent,
            "qty":      qty
        })

    return result


def parse_multipart(body, boundary):
    parts = {}
    boundary_bytes = ("--" + boundary).encode()
    chunks = body.split(boundary_bytes)
    for chunk in chunks:
        if b"Content-Disposition" not in chunk:
            continue
        header_end = chunk.find(b"\r\n\r\n")
        if header_end == -1:
            continue
        header = chunk[:header_end].decode("utf-8", errors="ignore")
        data   = chunk[header_end + 4:]
        if data.endswith(b"\r\n"):
            data = data[:-2]
        name = ""
        for part in header.split(";"):
            part = part.strip()
            if part.startswith("name="):
                name = part[5:].strip('"')
        if name:
            parts[name] = data
    return parts


class handler(BaseHTTPRequestHandler):

    def do_OPTIONS(self):
        self.send_response(200)
        self._cors()
        self.end_headers()

    def do_POST(self):
        try:
            content_type = self.headers.get("Content-Type", "")
            if "multipart/form-data" not in content_type:
                self._error(400, "multipart/form-data expected")
                return

            boundary = ""
            for part in content_type.split(";"):
                part = part.strip()
                if part.startswith("boundary="):
                    boundary = part[9:].strip('"')

            length = int(self.headers.get("Content-Length", 0))
            body   = self.rfile.read(length)
            parts  = parse_multipart(body, boundary)

            if "file" not in parts:
                self._error(400, "file required")
                return

            wb   = openpyxl.load_workbook(io.BytesIO(parts["file"]))
            rows = read_raw_bom(wb)
            result = build_quantified_bom(rows, convert_revision=False)

            response = {"success": True, "rows": result, "count": len(result)}

            self.send_response(200)
            self._cors()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps(response).encode())

        except Exception as e:
            self._error(500, str(e))

    def _cors(self):
        self.send_header("Access-Control-Allow-Origin",  "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def _error(self, code, msg):
        self.send_response(code)
        self._cors()
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"success": False, "error": msg}).encode())

    def log_message(self, format, *args):
        pass
