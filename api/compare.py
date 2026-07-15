from http.server import BaseHTTPRequestHandler
import json
import io
import openpyxl

def read_bom(wb):
    ws = wb.worksheets[0]
    rows = []
    for r in range(2, ws.max_row + 1):
        r0 = ws.cell(r, 1).value
        if r0 is None:
            continue
        try:
            level = int(float(str(r0)))
        except:
            continue
        title      = str(ws.cell(r, 2).value or "").strip()
        rev        = str(ws.cell(r, 3).value or "").strip()
        parent     = str(ws.cell(r, 5).value or "").strip()
        qty_val    = ws.cell(r, 7).value
        qty        = int(float(str(qty_val))) if qty_val is not None else 0
        rows.append({
            "level": level, "title": title, "rev": rev,
            "parent": parent, "qty": qty
        })
    return rows


def compare_boms(old_bom, new_bom):
    old_matched = [False] * len(old_bom)
    results = []

    for n in new_bom:
        found_idx = -1
        for i, o in enumerate(old_bom):
            if old_matched[i]:
                continue
            if o["title"] == n["title"] and o["parent"] == n["parent"]:
                found_idx = i
                old_matched[i] = True
                break

        if found_idx == -1:
            status   = "ADDED"
            old_rev  = ""
            old_qty  = 0
        else:
            o       = old_bom[found_idx]
            old_rev = o["rev"]
            old_qty = o["qty"]
            if n["qty"] == 0:
                status = "REMOVED"
            else:
                parts = []
                if old_rev != n["rev"]:
                    parts.append("REVISION CHANGE")
                if old_qty != n["qty"]:
                    parts.append("QTY CHANGE")
                status = "\n".join(parts) if parts else "NOT CHANGE"

        results.append({
            "status":  status,
            "level":   n["level"],
            "title":   n["title"],
            "parent":  n["parent"],
            "oldRev":  old_rev,
            "newRev":  n["rev"],
            "oldQty":  old_qty,
            "newQty":  n["qty"]
        })

    # REMOVED: OLD BOM'da eslesmemis satirlar
    for i, o in enumerate(old_bom):
        if old_matched[i]:
            continue
        insert_after = -1
        for j, res in enumerate(results):
            if res["title"] == o["parent"]:
                insert_after = j
            if res["parent"] == o["parent"]:
                insert_after = j

        rm = {
            "status":  "REMOVED",
            "level":   o["level"],
            "title":   o["title"],
            "parent":  o["parent"],
            "oldRev":  o["rev"],
            "newRev":  "",
            "oldQty":  o["qty"],
            "newQty":  0
        }
        if insert_after == -1:
            results.append(rm)
        else:
            results.insert(insert_after + 1, rm)

    # OLD TITLES
    seen = set()
    old_titles = []
    for r in results:
        if ("REVISION CHANGE" in r["status"] or r["status"] == "REMOVED") and r["oldRev"]:
            key = (r["title"], r["oldRev"])
            if key not in seen:
                seen.add(key)
                old_titles.append({"title": r["title"], "oldRev": r["oldRev"]})

    return results, old_titles


def parse_multipart(body, boundary):
    """Multipart form-data parser - iki dosyayi ayiklar"""
    parts = {}
    boundary_bytes = ("--" + boundary).encode()
    chunks = body.split(boundary_bytes)
    for chunk in chunks:
        if b"Content-Disposition" not in chunk:
            continue
        header_end = chunk.find(b"\r\n\r\n")
        if header_end == -1:
            continue
        header = chunk[:header_end].decode("utf-8", errors="ignore")
        data   = chunk[header_end + 4:]
        if data.endswith(b"\r\n"):
            data = data[:-2]
        # name ayikla
        name = ""
        for part in header.split(";"):
            part = part.strip()
            if part.startswith("name="):
                name = part[5:].strip('"')
        if name:
            parts[name] = data
    return parts


class handler(BaseHTTPRequestHandler):

    def do_OPTIONS(self):
        self.send_response(200)
        self._cors()
        self.end_headers()

    def do_POST(self):
        try:
            content_type = self.headers.get("Content-Type", "")
            if "multipart/form-data" not in content_type:
                self._error(400, "multipart/form-data expected")
                return

            boundary = ""
            for part in content_type.split(";"):
                part = part.strip()
                if part.startswith("boundary="):
                    boundary = part[9:].strip('"')

            length = int(self.headers.get("Content-Length", 0))
            body   = self.rfile.read(length)

            parts  = parse_multipart(body, boundary)

            if "old" not in parts or "new" not in parts:
                self._error(400, "old and new files required")
                return

            wb_old = openpyxl.load_workbook(io.BytesIO(parts["old"]))
            wb_new = openpyxl.load_workbook(io.BytesIO(parts["new"]))

            old_bom = read_bom(wb_old)
            new_bom = read_bom(wb_new)

            results, old_titles = compare_boms(old_bom, new_bom)

            # Istatistik
            stats = {"added": 0, "removed": 0, "changed": 0, "notChange": 0}
            for r in results:
                if r["status"] == "ADDED":           stats["added"]     += 1
                elif r["status"] == "REMOVED":       stats["removed"]   += 1
                elif r["status"] == "NOT CHANGE":    stats["notChange"] += 1
                else:                                stats["changed"]   += 1

            response = {
                "success":   True,
                "results":   results,
                "oldTitles": old_titles,
                "stats":     stats
            }

            self.send_response(200)
            self._cors()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps(response).encode())

        except Exception as e:
            self._error(500, str(e))

    def _cors(self):
        self.send_header("Access-Control-Allow-Origin",  "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def _error(self, code, msg):
        self.send_response(code)
        self._cors()
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"success": False, "error": msg}).encode())

    def log_message(self, format, *args):
        pass
